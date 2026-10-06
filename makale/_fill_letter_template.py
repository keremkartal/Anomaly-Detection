"""
Hocanin verdigi sablonu YERINDE doldurur -- bicimlendirmesini koruyarak.

`_build_letter_docx.py` mektubu sifirdan kuruyordu, yani kendi duzenimde.
Bu script farkli: `yeniveriseti/Response_to_Reviewers_sensors-4582518.docx`
dosyasini acar, yalnizca doldurulmasi gereken paragraflari degistirir ve
geri kalan her seyi oldugu gibi birakir.

SABLONUN BICIMI (olculdu, varsayilmadi)
  "Comment N.M"                 koyu, 11 pt, renk 1F3864
  hakem yorumu                  girinti 0.25", golge F2F2F2, italik, 404040
  "Response: " + govde         prefix koyu
  "Changes in the manuscript: " prefix koyu

NE DEGISIR
  - 0. paragraf (NOTE FOR AUTHORS) silinir
  - "Summary of main changes" maddeleri gercek sonuclara gore yazilir
  - her yorumun Response blogu HAKEM_YANIT.md'deki cevapla degisir
  - "Changes in the manuscript:" satiri cevabin ICINDEKI gondermelerden
    turetilir (Section/Table/Figure/Appendix/Equation), sirasi korunarak

NE DEGISMEZ
  - hakem yorumlarinin metni ve gri kutulari
  - basliklar, stiller, punto, renkler, paragraf sirasi

KULLANIM
    python makale/_fill_letter_template.py [--out YOL]
"""
import argparse
import copy
import pathlib
import re
import sys

sys.stdout.reconfigure(encoding="utf-8")

HERE = pathlib.Path(__file__).parent
ROOT = HERE.parent
TEMPLATE = ROOT / "yeniveriseti" / "Response_to_Reviewers_sensors-4582518.docx"
LETTER_MD = ROOT / "HAKEM_YANIT.md"
DEFAULT_OUT = ROOT / "Response_to_Reviewers_sensors-4582518_DOLDURULMUS.docx"

ITEM_RE = re.compile(r"^###\s+([\d.,\s]+?)\s*—\s*(.*)$")
LOC_RE = re.compile(
    r"\b(?:Section~?|Sections~?)\s*([IVX]+(?:-[A-Z]\d?)?)"
    r"|\bTable\s*(\d+)"
    r"|\bFigure\s*(\d+)"
    r"|\bAppendix\s*([A-C])"
    r"|\bEquation\s*(\d+)")

# Mektup numara TASIMAZ; {tab:cv} gibi etiket tasir ve numarasi buradan,
# derlenmis aux'tan cozulur. Dokuz gonderme bir kez yanlis tabloyu
# gosterdigi icin -- numaralari elle yazmistim ve sonradan tablo eklenince
# kaymislardi -- artik elle numara yazilmiyor.
AUX = ROOT / "makale" / "main_revised.aux"
LABEL_RE = re.compile(r"\{((?:tab|fig|eq):[a-z_0-9]+)\}")
_AUX_PAT = (r"newlabel\{((?:tab|fig|eq):[a-z_0-9]+)\}\{\{(?:"
            + "\\\\" + r"mbox\s*\{)?([0-9]+)")
KIND = {"tab": "Table", "fig": "Figure", "eq": "Equation"}


def label_numbers():
    if not AUX.exists():
        return {}
    aux = AUX.read_text(encoding="utf-8", errors="ignore")
    return {m.group(1): m.group(2) for m in re.finditer(_AUX_PAT, aux)}


def resolve(text, numbers, unresolved):
    def rep(m):
        lbl = m.group(1)
        n = numbers.get(lbl)
        if not n:
            unresolved.add(lbl)
            return "{" + lbl + "}"
        return "%s %s" % (KIND[lbl.split(":")[0]], n)
    return LABEL_RE.sub(rep, text)

# Sonuclara gore yeniden yazilan ozet maddeleri.
SUMMARY = [
    "Cross-validated threshold procedure corrected: each fold-seed run's test "
    "predictions are thresholded only with that run's own validation-selected "
    "threshold. Every threshold-dependent table was recomputed.",
    "Graph construction is no longer selected with label information. Both "
    "constructions are prespecified factors, edge generation and access to "
    "held-out topology are documented, and frequency-based node statistics "
    "are computed from training trips only.",
    "Cross-validation extended from two to five seeds (42-46), giving 25 runs "
    "per configuration and 125 in total, and including seed 45. Extending it "
    "changed the fusion ranking: the gated operator is first of four over two "
    "seeds and fourth over five, and the temporal-only control is first.",
    "Uncertainty estimation now accounts for seed variability and overlapping "
    "training sets through the Nadeau-Bengio correction, which widens "
    "intervals by a factor of 4.03 here. Window- and trip-level tests use the "
    "same F1 statistic and the same Holm correction, so the only difference "
    "between them is the statistical unit.",
    "Gradient-boosting baselines and the three-threshold rule added to "
    "grouped cross-validation and to the trip-level analysis. They are ahead "
    "of the hybrid: LightGBM reaches 0.9195 against 0.8589, a trip-level "
    "difference of +0.0409 with Holm p = 0.0002. This is the only performance "
    "difference the study can resolve, and it runs against the proposed "
    "architecture.",
    "New control experiments (edge removal, node-feature shuffling, random "
    "topology of matched density, non-graph contextual features), sensitivity "
    "analyses (window and stride; four treatments of the imputed attributes; "
    "anomaly prevalence; sensor noise), an inductive check, and epoch-level "
    "gate and gradient diagnostics. None of the graph controls differs from "
    "the reference: every Holm-corrected p is 1.000.",
    "Subtype analysis extended to all subtypes present in the data, with "
    "explicit definitions and precision, recall and false-positive counts, "
    "and with sparse subtypes reported descriptively only.",
    "Language revised throughout. Claims that a non-significant result shows "
    "absence of effect were removed; where equivalence is asserted it is "
    "supported by a TOST test with a margin fixed in advance. Three "
    "appendices were added, and a replication across four independently "
    "generated datasets spanning two road networks leaves 0 of 40 pairwise "
    "comparisons significant.",
]


def parse_answers():
    text = LETTER_MD.read_text(encoding="utf-8")
    out, cur, buf = {}, None, []

    def flush():
        if cur:
            body = "\n".join(buf).strip()
            for num in cur:
                out[num] = body

    for ln in text.split("\n"):
        m = ITEM_RE.match(ln)
        if m:
            flush()
            cur = [n.strip() for n in m.group(1).replace(",", " ").split()]
            buf = []
        elif cur is not None:
            if ln.startswith("## "):
                flush()
                cur, buf = None, []
            else:
                buf.append(ln)
    flush()
    return out


def locations(body):
    seen, out = set(), []
    for m in LOC_RE.finditer(body):
        sec, tab, fig, app, eq = m.groups()
        if sec:
            lbl = f"Section {sec}"
        elif tab:
            lbl = f"Table {tab}"
        elif fig:
            lbl = f"Figure {fig}"
        elif app:
            lbl = f"Appendix {app}"
        else:
            lbl = f"Equation {eq}"
        if lbl not in seen:
            seen.add(lbl)
            out.append(lbl)
    return out


def flatten(body):
    """Markdown govdeyi paragraf listesine cevirir; tablolari satira indirir."""
    paras = []
    for block in body.split("\n\n"):
        block = block.strip()
        if not block or set(block) <= set("-= "):
            continue
        if block.startswith("|"):
            rows = [r.strip().strip("|").split("|")
                    for r in block.split("\n") if r.strip().startswith("|")]
            rows = [[c.strip() for c in r] for r in rows]
            rows = [r for r in rows if not all(set(c) <= set("-: ") for c in r)]
            if len(rows) >= 2:
                head = rows[0]
                for r in rows[1:]:
                    paras.append("    " + "; ".join(
                        f"{h} {v}" for h, v in zip(head, r) if v))
            continue
        paras.append(" ".join(block.split("\n")))
    return paras


def set_text(par, segments, template_run):
    """Paragrafin run'larini silip yenilerini sablon run bicimiyle yazar."""
    for r in list(par.runs):
        r._r.getparent().remove(r._r)
    for text, bold in segments:
        if not text:
            continue
        r = par.add_run(text)
        r.bold = bold
        r.italic = template_run.italic
        if template_run.font.size:
            r.font.size = template_run.font.size
        if template_run.font.name:
            r.font.name = template_run.font.name
    return par


def md_segments(text):
    segs = []
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if part:
            segs.append((part.replace("*", ""), bool(i % 2)))
    return segs


def fill(out_path):
    from docx import Document

    answers = parse_answers()
    numbers = label_numbers()
    unresolved = set()
    answers = {k: resolve(v, numbers, unresolved) for k, v in answers.items()}
    doc = Document(str(TEMPLATE))
    body = doc.element.body

    # --- NOTE FOR AUTHORS paragrafini sil
    removed_note = False
    for p in doc.paragraphs:
        if p.text.strip().startswith("NOTE FOR AUTHORS"):
            p._p.getparent().remove(p._p)
            removed_note = True
            break

    # --- ozet maddeleri
    n_sum = 0
    for p in doc.paragraphs:
        m = re.match(r"^(\d)\.\s", p.text.strip())
        if m and p.text.strip()[0].isdigit() and len(p.runs) and n_sum < len(SUMMARY):
            idx = int(m.group(1))
            if 1 <= idx <= len(SUMMARY) and idx == n_sum + 1:
                set_text(p, [(f"{idx}. ", True),
                             (SUMMARY[idx - 1], False)], p.runs[0])
                n_sum += 1

    # --- yorum bloklari
    paras = doc.paragraphs
    starts = [(i, re.match(r"^Comment\s+(\d+\.\d+)$", p.text.strip()).group(1))
              for i, p in enumerate(paras)
              if re.match(r"^Comment\s+(\d+\.\d+)$", p.text.strip())]

    n_resp = n_chg = n_miss = 0
    for k, (i, num) in enumerate(starts):
        end = starts[k + 1][0] if k + 1 < len(starts) else len(paras)
        ans = answers.get(num)
        if not ans:
            n_miss += 1
            continue

        resp_idx = [j for j in range(i, end)
                    if paras[j].text.strip().startswith("Response:")]
        chg_idx = [j for j in range(i, end)
                   if paras[j].text.strip().startswith("Changes in the manuscript:")]
        if not resp_idx:
            n_miss += 1
            continue

        first = resp_idx[0]
        tmpl_run = paras[first].runs[0]
        blocks = flatten(ans)

        # ilk paragraf: "Response: " + ilk blok
        set_text(paras[first], [("Response: ", True)] + md_segments(blocks[0]),
                 tmpl_run)

        # Response ile Changes ARASINDAKI her sablon paragrafini sil.
        # (Yalnizca "Response:" ile baslayanlari silmek yetmiyor: sablonun
        #  devam paragraflari prefix tasimadigi icin ayakta kaliyor ve
        #  doldurulmamis parantezleri beraberinde getiriyor.)
        stop = chg_idx[0] if chg_idx else end
        for j in range(first + 1, stop):
            paras[j]._p.getparent().remove(paras[j]._p)

        # kalan bloklari ilk paragrafin ardina ekle
        anchor = paras[first]._p
        for blk in blocks[1:]:
            new = copy.deepcopy(paras[first]._p)
            anchor.addnext(new)
            anchor = new
        # yeni eklenenleri doldur
        from docx.text.paragraph import Paragraph
        cur = paras[first]._p
        for blk in blocks[1:]:
            cur = cur.getnext()
            set_text(Paragraph(cur, paras[first]._parent),
                     md_segments(blk), tmpl_run)
        n_resp += 1

        # --- Changes in the manuscript
        if chg_idx:
            cp = paras[chg_idx[0]]
            locs = locations(ans)
            txt = "; ".join(locs) + "." if locs else "see the response above."
            set_text(cp, [("Changes in the manuscript: ", True), (txt, False)],
                     cp.runs[0])
            n_chg += 1

    # --- sablonun kendi iki cumlesi
    # (a) sablon satir numarasi vaat ediyor; bizim gondermelerimiz bolum ve
    #     tablo numarasi, dizgiden bagimsiz. Cumle bunu soylemeli.
    # (b) sablon uydurma bir DOI ile bitiyor. Kavram DOI'si GERCEK ve kalici;
    #     surum DOI'si release kesilince eklenecek. DOI uydurulmaz.
    for p in doc.paragraphs:
        t = p.text.strip()
        if t.startswith("Below, reviewer comments are reproduced") and p.runs:
            set_text(p, [
                ("Below, reviewer comments are reproduced in grey boxes, "
                 "followed by our response and the corresponding changes. "
                 "References point to numbered sections, tables, figures and "
                 "appendices of the revised manuscript rather than to line "
                 "numbers, so that they remain valid under the journal's "
                 "typesetting. All changes are highlighted ", False),
                ("[specify: in blue / with tracked changes]", True),
                (".", False)], p.runs[0])
        elif t.startswith("We believe these revisions have substantially") and p.runs:
            set_text(p, [
                ("We believe these revisions have substantially strengthened "
                 "the manuscript, and we thank the reviewers again for their "
                 "time. The code, data, and result files of the revised "
                 "version, together with the scripts that verify every number "
                 "in the manuscript and in this letter against the stored "
                 "result files, are archived under the concept identifier ",
                 False),
                ("doi:10.5281/zenodo.22637360", True),
                (", which resolves to the most recent version. The version "
                 "identifier for the revised snapshot (v3.0.0) is added once "
                 "that release is deposited.", False)], p.runs[0])

    doc.save(str(out_path))
    return dict(note=removed_note, summary=n_sum, responses=n_resp,
                changes=n_chg, missing=n_miss, comments=len(starts),
                labels=len(numbers), unresolved=sorted(unresolved))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=str(DEFAULT_OUT))
    a = ap.parse_args()
    for f in (TEMPLATE, LETTER_MD):
        if not f.exists():
            sys.exit(f"yok: {f}")
    st = fill(a.out)
    print(f"yazildi: {a.out}\n")
    print(f"  sablondaki yorum        : {st['comments']}")
    print(f"  doldurulan Response     : {st['responses']}")
    print(f"  yazilan Changes satiri  : {st['changes']}")
    print(f"  yeniden yazilan ozet    : {st['summary']} / {len(SUMMARY)}")
    print(f"  NOTE paragrafi silindi  : {st['note']}")
    print(f"  aux'tan okunan etiket   : {st['labels']}")
    if st["missing"]:
        print(f"  CEVABI BULUNAMAYAN      : {st['missing']}")
    if st["unresolved"]:
        print()
        print(f"  COZULEMEYEN ETIKET      : {len(st['unresolved'])}")
        for lbl in st["unresolved"]:
            print(f"      {lbl}")
        sys.exit(1)


if __name__ == "__main__":
    main()
