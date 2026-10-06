"""
Yanit mektubunu MDPI'ya gidecek .docx olarak uretir.

NICIN BU SCRIPT VAR
-------------------
Resmi teslim bicimi `Response_to_Reviewers_sensors-4582518.docx`. Icerik
`HAKEM_YANIT.md`'de; bu script ikisini birlestirir: sablonun hakem
yorumlarini kendi dosyasindan okur, cevaplari Markdown'dan alir, ve
doldurulmus mektubu yazar.

Elle kopyalamak yerine script olmasinin sebebi: cevaplar hala degisiyor
(bes tohumlu kosular yedi sayiyi degistirdi) ve her degisiklikte 26 maddeyi
elle tasimak hata kaynagi. Dogrulayici `_verify_letter.py` Markdown'i
denetliyor; bu script de denetlenmis Markdown'dan uretiyor.

GIZLILIK
--------
Cikti `*.docx` ile gitignore'da. Hakem yorumlari depoya girmez.

KULLANIM
--------
    python makale/_build_letter_docx.py [--out YOL]
"""
import argparse
import html
import pathlib
import re
import sys
import zipfile

sys.stdout.reconfigure(encoding="utf-8")

HERE = pathlib.Path(__file__).parent
ROOT = HERE.parent
LETTER_MD = ROOT / "HAKEM_YANIT.md"
TEMPLATE = ROOT / "yeniveriseti" / "Response_to_Reviewers_sensors-4582518.docx"
DEFAULT_OUT = ROOT / "Response_to_Reviewers_sensors-4582518_FILLED.docx"

# Markdown basligindaki madde numarasi -> sablondaki "Comment N.M"
ITEM_RE = re.compile(r"^###\s+([\d.,\s]+?)\s*—\s*(.*)$")


def template_comments():
    """Sablondan hakem yorumlarini madde numarasina gore cikarir."""
    if not TEMPLATE.exists():
        return {}
    z = zipfile.ZipFile(TEMPLATE)
    x = z.read("word/document.xml").decode("utf-8")
    x = re.sub(r"</w:p>", "\n", x)
    x = re.sub(r"<[^>]+>", "", x)
    lines = [html.unescape(l).strip() for l in x.split("\n") if html.unescape(l).strip()]
    out, cur = {}, None
    for ln in lines:
        m = re.match(r"^Comment\s+(\d+\.\d+)$", ln)
        if m:
            cur = m.group(1)
            out[cur] = []
            continue
        if cur:
            if ln.startswith(("Response:", "Changes in the manuscript:",
                              "Comment ", "Response to Reviewer")):
                cur = None
                continue
            out[cur].append(ln)
    return {k: " ".join(v) for k, v in out.items() if v}


def parse_answers():
    """HAKEM_YANIT.md'yi madde bloklarina ayirir."""
    text = LETTER_MD.read_text(encoding="utf-8")
    blocks, cur, buf = [], None, []
    for ln in text.split("\n"):
        m = ITEM_RE.match(ln)
        if m:
            if cur:
                blocks.append((cur[0], cur[1], "\n".join(buf).strip()))
            nums = [n.strip() for n in m.group(1).replace(",", " ").split()]
            cur, buf = (nums, m.group(2)), []
        elif cur is not None:
            if ln.startswith("## "):
                blocks.append((cur[0], cur[1], "\n".join(buf).strip()))
                cur, buf = None, []
            else:
                buf.append(ln)
    if cur:
        blocks.append((cur[0], cur[1], "\n".join(buf).strip()))
    return blocks


def md_runs(par, text):
    """**kalin** isaretlemesini docx run'larina cevirir."""
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if not part:
            continue
        r = par.add_run(part)
        r.bold = bool(i % 2)


def build(out_path):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    answers = parse_answers()
    comments = template_comments()
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)

    h = doc.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = h.add_run("Response to Reviewers")
    r.bold = True
    r.font.size = Pt(16)

    for line in ("Manuscript ID: sensors-4582518",
                 "How Much Does Fusion Design Matter? Partition, Seed, and "
                 "Cluster-Aware Evaluation of Hybrid LSTM–GAT Traffic "
                 "Anomaly Detectors"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rr = p.add_run(line)
        rr.italic = True
    doc.add_paragraph()

    # giris: Markdown'in ilk bolumu (madde basliklarindan once)
    intro = LETTER_MD.read_text(encoding="utf-8").split("## Reviewer 1")[0]
    for para in intro.split("\n\n"):
        para = para.strip()
        if not para or para.startswith("#") or set(para) <= set("-= "):
            continue
        md_runs(doc.add_paragraph(), " ".join(para.split("\n")))

    cur_rev = None
    n_items = n_comments = 0
    for nums, title, body in answers:
        rev = nums[0].split(".")[0]
        if rev != cur_rev:
            cur_rev = rev
            doc.add_paragraph()
            p = doc.add_paragraph()
            rr = p.add_run(f"Response to Reviewer {rev}")
            rr.bold = True
            rr.font.size = Pt(13)

        p = doc.add_paragraph()
        rr = p.add_run("Comment " + ", ".join(nums))
        rr.bold = True
        n_items += len(nums)

        # sablondaki hakem yorumu, varsa
        got = [comments[n] for n in nums if n in comments]
        if got:
            n_comments += len(got)
            for g in got:
                q = doc.add_paragraph()
                q.paragraph_format.left_indent = Pt(18)
                rq = q.add_run(g)
                rq.italic = True
                rq.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

        p = doc.add_paragraph()
        rr = p.add_run(title)
        rr.bold = True

        for para in body.split("\n\n"):
            para = para.strip()
            if not para or set(para) <= set("-= "):
                continue
            if para.startswith("|"):          # Markdown tablosu -> docx tablo
                rows = [r for r in para.split("\n") if r.strip().startswith("|")]
                cells = [[c.strip() for c in r.strip().strip("|").split("|")]
                         for r in rows]
                cells = [c for c in cells if not all(set(x) <= set("-: ") for x in c)]
                if not cells:
                    continue
                t = doc.add_table(rows=0, cols=len(cells[0]))
                t.style = "Light Grid Accent 1"
                for ri, row in enumerate(cells):
                    dc = t.add_row().cells
                    for ci, val in enumerate(row[:len(cells[0])]):
                        dc[ci].text = ""
                        md_runs(dc[ci].paragraphs[0], val)
                        if ri == 0:
                            for rn in dc[ci].paragraphs[0].runs:
                                rn.bold = True
                continue
            md_runs(doc.add_paragraph(), " ".join(para.split("\n")))

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.add_run("Sincerely,").italic = True
    doc.add_paragraph("Ahmet Sayar, on behalf of all authors")

    doc.save(out_path)
    return len(answers), n_items, n_comments


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=str(DEFAULT_OUT))
    a = ap.parse_args()
    if not LETTER_MD.exists():
        sys.exit(f"mektup yok: {LETTER_MD}")
    nb, ni, nc = build(a.out)
    print(f"yazildi: {a.out}")
    print(f"  madde blogu      : {nb}")
    print(f"  kapsanan madde   : {ni}")
    print(f"  sablondan alinan hakem yorumu : {nc}")
    if nc == 0:
        print("  NOT: sablon bulunamadi, yorumlar bos birakildi")


if __name__ == "__main__":
    main()
